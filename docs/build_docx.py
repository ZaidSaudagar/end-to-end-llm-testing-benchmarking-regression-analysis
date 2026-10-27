"""Convert docs/PROJECT_DOCUMENTATION.md into a formatted Word document.

Produces real Word tables (not flattened text), styled headings, a table of
contents field, and embedded Mermaid diagram source blocks.

Usage:
    .venv\\Scripts\\python.exe docs\\build_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCS_DIR = Path(__file__).resolve().parent
SOURCE = DOCS_DIR / "PROJECT_DOCUMENTATION.md"
TARGET = DOCS_DIR / "PROJECT_DOCUMENTATION.docx"
DIAGRAMS_DIR = DOCS_DIR / "diagrams"

ACCENT = RGBColor(0xF4, 0x3F, 0x5E)
INK = RGBColor(0x1A, 0x1C, 0x20)
MUTED = RGBColor(0x5A, 0x61, 0x6E)
CODE_BG = "F4F5F7"
HEADER_BG = "1A1C20"


# --------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------


def shade(cell, hex_fill: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_rule(paragraph, color: str = "D8DBE0", size: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


# --------------------------------------------------------------------------
# inline markdown
# --------------------------------------------------------------------------

INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*|!\[[^\]]*\]\([^)]+\)|\[[^\]]+\]\([^)]+\))",
    re.DOTALL,
)


def add_inline(paragraph, text: str, *, base_size: float = 10.5, bold_all: bool = False) -> None:
    """Render inline markdown (bold, code, italic, links) into a paragraph."""
    for token in INLINE_RE.split(text):
        if not token:
            continue

        if token.startswith("!["):
            # images are handled at block level; skip stray inline occurrences
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 0.5)
            run.font.color.rgb = RGBColor(0xB0, 0x2A, 0x45)
        elif token.startswith("[") and "](" in token:
            label, _, url = token[1:-1].partition("](")
            run = paragraph.add_run(label)
            run.font.color.rgb = ACCENT
            run.underline = True
            if url and not url.startswith("#"):
                run.font.size = Pt(base_size)
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token)

        if run.font.size is None:
            run.font.size = Pt(base_size)
        if bold_all:
            run.bold = True


# --------------------------------------------------------------------------
# block parsing
# --------------------------------------------------------------------------


def parse_blocks(md: str) -> list[tuple[str, object]]:
    """Split markdown into (kind, payload) blocks."""
    lines = md.split("\n")
    blocks: list[tuple[str, object]] = []
    i = 0

    while i < len(lines):
        line = lines[i]

        # fenced code
        if line.strip().startswith("```"):
            lang = line.strip()[3:].strip()
            i += 1
            buf: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", (lang, "\n".join(buf))))
            continue

        # table
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(
            r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]
        ):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(("table", (header, rows)))
            continue

        # image
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", line.strip())
        if m:
            blocks.append(("image", (m.group(1), m.group(2))))
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            blocks.append(("heading", (len(m.group(1)), m.group(2).strip())))
            i += 1
            continue

        # horizontal rule
        if re.match(r"^\s*---+\s*$", line):
            blocks.append(("rule", None))
            i += 1
            continue

        # blockquote
        if line.strip().startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(buf)))
            continue

        # list
        if re.match(r"^\s*([-*]|\d+\.)\s+", line):
            ordered = bool(re.match(r"^\s*\d+\.\s+", line))
            items: list[str] = []
            while i < len(lines) and re.match(r"^\s*([-*]|\d+\.)\s+", lines[i]):
                items.append(re.sub(r"^\s*([-*]|\d+\.)\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("list", (ordered, items)))
            continue

        # blank
        if not line.strip():
            i += 1
            continue

        # paragraph
        buf = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^\s*(#{1,6}\s|[-*]\s|\d+\.\s|\||>|```|---+\s*$)", lines[i]
        ):
            buf.append(lines[i].strip())
            i += 1
        blocks.append(("para", " ".join(buf)))

    return blocks


# --------------------------------------------------------------------------
# rendering
# --------------------------------------------------------------------------


def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in (
        ("Heading 1", 20, INK, 22, 8),
        ("Heading 2", 15, INK, 18, 6),
        ("Heading 3", 12.5, INK, 14, 4),
        ("Heading 4", 11, MUTED, 10, 3),
    ):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("End-to-End LLM Testing,\nBenchmarking and Regression Analysis")
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = INK
    title.paragraph_format.space_after = Pt(6)

    rule = doc.add_paragraph()
    add_rule(rule, color="F43F5E", size=18)
    rule.paragraph_format.space_after = Pt(14)

    sub = doc.add_paragraph()
    r = sub.add_run("Project Documentation")
    r.font.size = Pt(15)
    r.font.color.rgb = MUTED
    sub.paragraph_format.space_after = Pt(28)

    for label, value in (
        ("Document", "PROJECT_DOCUMENTATION.md / .docx"),
        ("Scope", "Architecture, modules, APIs, database, design system, testing"),
        ("Backend", "Python 3.11+ · FastAPI · SQLAlchemy 2.0 · SQLite"),
        ("Frontend", "React 19 · TypeScript · Vite 8"),
        ("Diagrams", "12 Mermaid sources in docs/diagrams/"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        rl = p.add_run(f"{label}   ")
        rl.bold = True
        rl.font.size = Pt(10)
        rl.font.color.rgb = INK
        rv = p.add_run(value)
        rv.font.size = Pt(10)
        rv.font.color.rgb = MUTED

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc: Document) -> None:
    h = doc.add_heading("Contents", level=1)
    h.paragraph_format.space_before = Pt(0)

    note = doc.add_paragraph()
    r = note.add_run(
        "Word will populate this field on open. If it appears empty, "
        "right-click it and choose Update Field."
    )
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = MUTED

    p = doc.add_paragraph()
    add_field(p, r'TOC \o "1-2" \h \z \u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True

    hdr = table.rows[0].cells
    for idx, text in enumerate(header):
        hdr[idx].text = ""
        p = hdr[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, text, base_size=9.5, bold_all=True)
        for run in p.runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(hdr[idx], HEADER_BG)
        set_cell_margins(hdr[idx])

    for row in rows:
        cells = table.add_row().cells
        for idx in range(len(header)):
            value = row[idx] if idx < len(row) else ""
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value, base_size=9.5)
            set_cell_margins(cells[idx])

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, lang: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    shade(cell, CODE_BG)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    cell.text = ""

    first = True
    for line in body.split("\n"):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)

    if lang:
        cap = doc.add_paragraph()
        cap.paragraph_format.space_before = Pt(2)
        r = cap.add_run(lang)
        r.font.size = Pt(8)
        r.italic = True
        r.font.color.rgb = MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_quote(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    shade(cell, "FDF2F4")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, text, base_size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc: Document, alt: str, rel_path: str) -> None:
    """Embed a PNG screenshot, or describe a Mermaid source file."""
    target = (DOCS_DIR / rel_path).resolve()

    if target.suffix.lower() == ".mmd":
        # Mermaid source cannot be rasterised here; render it as a labelled
        # code block so the diagram definition travels with the document.
        if not target.exists():
            return
        body = target.read_text(encoding="utf-8")
        cap = doc.add_paragraph()
        cap.paragraph_format.space_after = Pt(2)
        r = cap.add_run(f"Diagram — {alt}")
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = INK
        add_code_block(doc, "mermaid", body)
        return

    if not target.exists():
        p = doc.add_paragraph()
        r = p.add_run(f"[missing image: {rel_path}]")
        r.italic = True
        r.font.color.rgb = MUTED
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    try:
        run.add_picture(str(target), width=Inches(6.4))
    except Exception:
        run.text = f"[unreadable image: {rel_path}]"

    if alt:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        r = cap.add_run(alt)
        r.italic = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = MUTED


def render(doc: Document, blocks: list[tuple[str, object]]) -> None:
    for kind, payload in blocks:
        if kind == "heading":
            level, text = payload  # type: ignore[misc]
            if level == 1:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(text, level=min(level, 4))

        elif kind == "para":
            p = doc.add_paragraph()
            add_inline(p, payload)  # type: ignore[arg-type]

        elif kind == "list":
            ordered, items = payload  # type: ignore[misc]
            for idx, item in enumerate(items, start=1):
                p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, item)

        elif kind == "table":
            header, rows = payload  # type: ignore[misc]
            add_table(doc, header, rows)

        elif kind == "code":
            lang, body = payload  # type: ignore[misc]
            add_code_block(doc, lang, body)

        elif kind == "quote":
            add_quote(doc, payload)  # type: ignore[arg-type]

        elif kind == "image":
            alt, rel_path = payload  # type: ignore[misc]
            add_image(doc, alt, rel_path)

        elif kind == "rule":
            p = doc.add_paragraph()
            add_rule(p)
            p.paragraph_format.space_after = Pt(6)


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("End-to-End LLM Testing, Benchmarking and Regression Analysis  ·  ")
        r.font.size = Pt(8)
        r.font.color.rgb = MUTED
        add_field(p, "PAGE")
        for run in p.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = MUTED


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Source not found: {SOURCE}")

    md = SOURCE.read_text(encoding="utf-8")
    blocks = parse_blocks(md)

    doc = Document()
    style_document(doc)
    add_cover(doc)
    add_toc(doc)
    render(doc, blocks)
    add_footer(doc)

    doc.save(TARGET)

    tables = len(doc.tables)
    paras = len(doc.paragraphs)
    print(f"Wrote {TARGET}")
    print(f"  blocks parsed : {len(blocks)}")
    print(f"  paragraphs    : {paras}")
    print(f"  tables        : {tables}")
    print(f"  size          : {TARGET.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
